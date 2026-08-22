from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "testing" / "results"
ASSETS = RESULTS / "report-assets"
OUTPUT = ROOT / "testing" / "reports" / "monolith-e2e-performance-baseline.docx"

# Document skill preset: standard_business_brief.
NAVY = "0F172A"
SLATE = "334155"
MUTED = "64748B"
BLUE = "2563EB"
CYAN = "0EA5E9"
GREEN = "059669"
AMBER = "D97706"
RED = "DC2626"
LIGHT_BLUE = "EFF6FF"
LIGHT_GREEN = "ECFDF5"
LIGHT_AMBER = "FFFBEB"
LIGHT_RED = "FEF2F2"
LIGHT = "F8FAFC"
LINE = "CBD5E1"
WHITE = "FFFFFF"


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


SUMMARY = load_json(RESULTS / "summary" / "monolith-summary.json")
PROFILE = load_json(RESULTS / "summary" / "profile-summary.json")
SYSTEM = load_json(RESULTS / "summary" / "system.json")
DB = load_json(RESULTS / "database" / "performance-audit.json")
LH = load_json(RESULTS / "lighthouse" / "scores.json")


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        for key, value in edges[edge].items():
            el.set(qn(f"w:{key}"), str(value))


def cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_column_widths(table, widths_cm):
    for row in table.rows:
        for index, width in enumerate(widths_cm):
            if index < len(row.cells):
                row.cells[index].width = Cm(width)


def add_field(run, instruction: str):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(SLATE)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for name, size, color in (("Title", 30, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 10.5, SLATE)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    styles["Heading 1"].paragraph_format.page_break_before = False
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(8)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
    styles["Caption"].paragraph_format.space_before = Pt(3)
    styles["Caption"].paragraph_format.space_after = Pt(7)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("DRIVEFLOW  /  MONOLITH BASELINE")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Bachelor’s thesis working report  •  19 August 2026  •  ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    page_run = p.add_run()
    page_run.font.size = Pt(7.5)
    add_field(page_run, "PAGE")

    props = doc.core_properties
    props.title = "DriveFlow Monolith — End-to-End and Performance Baseline"
    props.subject = "Reproducible baseline for monolith-to-microservices comparison"
    props.author = "DriveFlow thesis project"
    props.keywords = "monolith, microservices, k6, Playwright, MySQL, Lighthouse, performance"


def add_text(doc, text: str, bold_lead: str | None = None, color: str | None = None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(color or NAVY)
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2.5)
    p.add_run(text)
    return p


def add_number(doc, number: int, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.space_after = Pt(3)
    marker = p.add_run(f"{number}.  ")
    marker.bold = True
    marker.font.color.rgb = RGBColor.from_string(MUTED)
    p.add_run(text)
    return p


def add_callout(doc, title: str, body: str, kind: str = "info"):
    palette = {
        "info": (LIGHT_BLUE, BLUE),
        "pass": (LIGHT_GREEN, GREEN),
        "warn": (LIGHT_AMBER, AMBER),
        "fail": (LIGHT_RED, RED),
    }
    fill, accent = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_column_widths(table, [16.8])
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell_margins(cell, 120, 180, 120, 130)
    set_cell_border(
        cell,
        top={"val": "nil"},
        bottom={"val": "nil"},
        left={"val": "single", "sz": "32", "color": accent},
        right={"val": "nil"},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cell, 90, 85, 90, 85)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(label))
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(values):
            cell = cells[i]
            if row_index % 2 == 1:
                shade(cell, LIGHT)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell, 75, 85, 75, 85)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            set_cell_border(
                cell,
                bottom={"val": "single", "sz": "4", "color": LINE},
                left={"val": "single", "sz": "2", "color": "E2E8F0"},
                right={"val": "single", "sz": "2", "color": "E2E8F0"},
            )
    if widths:
        set_column_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_metric_cards(doc, cards):
    table = doc.add_table(rows=1, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (value, label, status) in enumerate(cards):
        cell = table.cell(0, i)
        fill = {"pass": LIGHT_GREEN, "fail": LIGHT_RED, "warn": LIGHT_AMBER, "info": LIGHT_BLUE}[status]
        accent = {"pass": GREEN, "fail": RED, "warn": AMBER, "info": BLUE}[status]
        shade(cell, fill)
        cell_margins(cell, 130, 90, 120, 90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(value)
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor.from_string(accent)
        p2 = cell.add_paragraph(label)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        p2.runs[0].font.size = Pt(7.5)
        p2.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, filename: str, caption: str, width_inches=6.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width_inches))
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def start_section(doc, number: str, title: str, kicker: str, page_break=True):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"SECTION {number}")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(CYAN)
    doc.add_heading(title, level=1)
    p = doc.add_paragraph(kicker)
    p.paragraph_format.space_after = Pt(9)
    r = p.runs[0]
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def fmt_ms(value):
    return f"{value:,.0f} ms" if value < 1000 else f"{value / 1000:,.2f} s"


def build_report():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # Cover — editorial band pattern applied to the standard business brief preset.
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner.autofit = False
    set_column_widths(banner, [17.2])
    cell = banner.cell(0, 0)
    shade(cell, NAVY)
    cell_margins(cell, 520, 380, 520, 380)
    p = cell.paragraphs[0]
    r = p.add_run("DRIVEFLOW  /  ARCHITECTURE BASELINE")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(CYAN)
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Monolith\nEnd-to-End &\nPerformance Baseline")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(29)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p = cell.add_paragraph("A reproducible evidence package for the bachelor’s thesis and the future microservices comparison")
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(0)
    p.runs[0].font.size = Pt(11)
    p.runs[0].font.color.rgb = RGBColor.from_string("CBD5E1")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_table(
        doc,
        ["Report status", "Measurement date", "System under test", "Dataset ceiling"],
        [["Working baseline · v1.0", "19 August 2026", "Node.js/Express + MySQL monolith", "50k customers · 75k leads"]],
        widths=[4.0, 3.4, 5.3, 4.1],
        font_size=8.5,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_callout(
        doc,
        "Purpose of this document",
        "Record the monolith’s observed correctness, capacity, failure modes, and resource behavior before architectural migration. The same automated suite is preserved as the comparison instrument for the microservices implementation.",
        "info",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Prepared from automated test artifacts\nE:\\thesis-project\\testing\\results")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    # Executive summary.
    start_section(doc, "01", "Executive summary", "The monolith is functionally complete under the tested flows, but its practical capacity is constrained by queued database work and CPU-heavy authentication.")
    add_metric_cards(
        doc,
        [
            ("6 / 6", "E2E workflows passed", "pass"),
            ("49 rps", "20-VU steady state", "pass"),
            ("0%", "read-workload failures", "pass"),
            ("1.30 s", "stress p95", "fail"),
        ],
    )
    doc.add_heading("Decision-grade findings", level=2)
    findings = [
        ("Functional readiness", "PASS", "Authentication, authorization, navigation, API contracts, and the complete dealership transaction flow passed in the production build."),
        ("Steady-state behavior", "PASS", "At 20 virtual users, the mixed read profile sustained 48.6 requests/s with 235.8 ms p95 latency and no request failures. The 30-minute soak reproduced this level without meaningful heap growth."),
        ("Capacity ceiling", "FAIL", "Throughput remained approximately 49–51 requests/s while concurrency grew from 20 to 200 users. p95 latency increased from 236 ms to 3.48 s, indicating queueing rather than added capacity."),
        ("Authentication", "FAIL", "The staged login test completed only 661 of 1,890 scheduled iterations, dropped 1,229, reached 40.43 s p95, and hit the 100-VU executor cap."),
        ("Large-data search", "FAIL", "Search/deep-page p95 grew from 18.8 ms at 1k customers to 1.23 s at 50k customers. EXPLAIN ANALYZE shows tenant-wide scans caused by CONCAT_WS(... ) LIKE '%term%'."),
        ("Transactional correctness", "MIXED", "Concurrent sale completion behaved correctly (1 success, 19 conflicts). Concurrent vehicle reservation preserved uniqueness, but 9 of 20 requests surfaced database deadlocks as HTTP 500 responses."),
    ]
    add_table(doc, ["Area", "Status", "Evidence"], findings, widths=[3.1, 1.8, 11.9], font_size=7.8)
    add_callout(
        doc,
        "Bottom line",
        "The monolith is suitable as a functional baseline. Before attributing improvements to microservices, the comparison must preserve data volume, request mix, environment, thresholds, and run order. Architecture alone should not be credited for changes caused by indexing, pool tuning, caching, or hardware.",
        "warn",
    )

    # Contents and scope.
    start_section(doc, "02", "Study design and scope", "A controlled, repeatable protocol designed for a monolith-to-microservices comparison.")
    doc.add_heading("Research framing", level=2)
    add_text(doc, "Independent variable: deployment architecture — the current monolith versus the future microservices implementation.", "Independent variable:")
    add_text(doc, "Dependent variables: correctness, throughput, p50/p95/p99 latency, error rate, dropped work, resource consumption, database behavior, and frontend quality.", "Dependent variables:")
    add_text(doc, "Controlled variables: workload scripts, dataset seeds, endpoint semantics, test host, database version and configuration, runtime version, thresholds, warm-up, and measurement duration.", "Controlled variables:")
    doc.add_heading("Test portfolio", level=2)
    add_table(
        doc,
        ["Layer", "Instrument", "Coverage", "Primary evidence"],
        [
            ("Functional UI/API", "Playwright", "Auth, refresh, logout, RBAC, routing, API errors, full sale workflow", "6 browser/API scenarios"),
            ("Performance", "Grafana k6 2.2.0", "Smoke, baseline, average, stress, spike, soak, breakpoint", "Latency, throughput, failures"),
            ("Scalability", "k6 + seeded datasets", "Search and deep pagination at 1k, 10k, and 50k customers", "Scale curve and EXPLAIN ANALYZE"),
            ("Correctness under race", "k6", "20-way reservation and 20-way sale completion", "Success/conflict/error distribution"),
            ("Runtime", "Node inspector + CSV monitor", "CPU, heap allocations, RSS, event loop", "CPU/heap profiles and 1 s telemetry"),
            ("Frontend quality", "Lighthouse 12.8.2", "Public login and authenticated dashboard", "Four category scores"),
        ],
        widths=[2.6, 3.2, 7.0, 4.0],
        font_size=7.7,
    )
    doc.add_heading("Acceptance thresholds used", level=2)
    add_bullet(doc, "Functional scenarios: 100% pass; authorization failures must return stable 401/403 responses.")
    add_bullet(doc, "Normal mixed reads: p95 < 500 ms and HTTP failure rate < 1%.")
    add_bullet(doc, "Large-data search guard: p95 < 750 ms.")
    add_bullet(doc, "Race tests: exactly one successful state transition; all losing attempts must return controlled conflicts, never HTTP 500.")
    add_bullet(doc, "Soak: stable request quality and no sustained RSS/heap growth trend over 30 minutes.")
    add_callout(doc, "Validity boundary", "Each performance profile was executed once on a local, shared test host with synthetic data. These measurements establish an engineering baseline, not population-level statistical confidence. For the final thesis, execute 3–5 repetitions per architecture and report median values plus dispersion or confidence intervals.", "warn")

    # Environment and datasets.
    start_section(doc, "03", "System under test", "Production frontend and backend builds, isolated performance database, and three deterministic dataset sizes.")
    add_table(
        doc,
        ["Component", "Measured configuration"],
        [
            ("Application", "Node.js v22.14.0 (x64), Express/TypeScript production build"),
            ("Database", f"MySQL {DB['database']['mysqlVersion']}; 128 MB InnoDB buffer pool; max_connections=151; Performance Schema ON"),
            ("Application DB pool", "connectionLimit=10; waitForConnections=true; queueLimit=0 (unbounded queue)"),
            ("Host", f"Windows {SYSTEM['release']}; {SYSTEM['cpuModel']}; {SYSTEM['logicalCpuCount']} logical CPUs; {SYSTEM['totalMemoryGb']:.2f} GB RAM"),
            ("Load topology", "Application, MySQL, browser, and load generator on the same host"),
            ("Transport", "Local HTTP; no TLS, WAN latency, reverse proxy, container scheduler, or autoscaling"),
        ],
        widths=[4.0, 12.8],
        font_size=8.2,
    )
    doc.add_heading("Seeded dataset sizes", level=2)
    add_table(
        doc,
        ["Dataset", "Vehicles", "Customers", "Leads", "Activities", "Offers", "Reservations", "Sales"],
        [
            ("Small", "500", "1,000", "1,500", "3,000", "500", "150", "200"),
            ("Medium", "5,000", "10,000", "15,000", "30,000", "5,000", "1,500", "2,000"),
            ("Large", "25,000", "50,000", "75,000", "150,000", "25,000", "7,500", "10,000"),
        ],
        widths=[2.2, 2.0, 2.2, 2.1, 2.3, 2.0, 2.2, 1.8],
        font_size=7.6,
    )
    add_text(doc, "The seed utility refuses to operate outside the dedicated performance database. Synthetic records preserve tenant scoping and realistic entity relationships while remaining reproducible.")
    doc.add_heading("Measurement caveats", level=2)
    add_bullet(doc, "Same-host resource contention can reduce absolute capacity and should be preserved or eliminated identically in both architecture runs.")
    add_bullet(doc, "The test uses generated data and a mostly warm local database cache; production distributions may differ.")
    add_bullet(doc, "Absolute throughput should be interpreted as host-specific. Relative behavior across load, scale, and architecture is the primary thesis value.")

    # E2E.
    start_section(doc, "04", "End-to-end correctness", "All six consolidated production-build scenarios passed in 14.9 seconds.")
    add_metric_cards(doc, [("6", "passed", "pass"), ("0", "failed", "pass"), ("14.9 s", "elapsed", "info"), ("1", "browser worker", "info")])
    add_table(
        doc,
        ["Scenario", "What it verified", "Result"],
        [
            ("API error contract", "401 unauthenticated, 400 invalid input, and 404 missing resource", "PASS"),
            ("Role enforcement", "Valid lower-privilege role receives 403 for a protected user endpoint", "PASS"),
            ("Invalid login", "Incorrect credentials are rejected without creating an authenticated session", "PASS"),
            ("Session lifecycle", "Login, protected dashboard, refresh/reload, and logout", "PASS"),
            ("Route protection", "Sidebar navigation, direct-route access, and role-based route denial", "PASS"),
            ("Dealership transaction", "Customer → vehicle → lead → activity → interest → test drive → offer → reservation → pending sale → completed sale", "PASS"),
        ],
        widths=[3.7, 11.2, 1.9],
        font_size=8.0,
    )
    add_callout(doc, "Interpretation", "The tested vertical slice is viable from UI to database and includes both happy paths and negative authorization/API behavior. This makes the suite suitable as a migration regression gate.", "pass")
    doc.add_heading("Coverage boundary", level=2)
    add_text(doc, "The suite is intentionally consolidated around the system’s most important workflow. It does not yet prove exhaustive field validation, cross-browser compatibility, accessibility interaction behavior, backup/restore, disaster recovery, or security penetration resistance.")

    # Performance overview.
    start_section(doc, "05", "Load, stress, spike, soak, and capacity", "Read-heavy workloads remained correct under pressure, but latency rose sharply after the 20-VU operating point.")
    k = SUMMARY["k6"]
    load_rows = []
    for key, label, status in [
        ("medium-baseline", "Baseline · 1→10 VUs", "PASS"),
        ("medium-average", "Average · 20 VUs / 10 min", "PASS"),
        ("medium-stress", "Stress · ramp to 100 VUs", "FAIL latency"),
        ("medium-spike", "Spike · 5→100 VUs", "FAIL latency"),
        ("medium-breakpoint", "Breakpoint · ramp to 200 VUs", "Observed limit"),
        ("medium-soak", "Soak · 20 VUs / 30 min", "PASS"),
    ]:
        m = k[key]
        load_rows.append((label, f"{m['requests']['count']:,}", f"{m['requests']['rate']:.1f}", fmt_ms(m['http']['avg']), fmt_ms(m['http']['p95']), fmt_ms(m['http']['p99']), "0%", status))
    add_table(
        doc,
        ["Profile", "Requests", "Req/s", "Average", "p95", "p99", "Failures", "Threshold"],
        load_rows,
        widths=[4.0, 1.8, 1.4, 1.8, 1.8, 1.8, 1.4, 2.8],
        font_size=7.2,
    )
    add_figure(doc, "load-profile-latency.png", "Figure 1. Tail latency crosses the 500 ms target under stress, spike, and breakpoint load. The logarithmic axis preserves both normal and overload values.")
    add_text(doc, "The spike profile’s 22.39 s p99 and 31.20 s maximum show a pronounced long tail. Zero read-workload HTTP failures should not be read as healthy service: the system continued accepting and queueing work while response time became operationally unusable.")

    doc.add_page_break()
    doc.add_heading("Capacity interpretation", level=1)
    add_figure(doc, "load-profile-throughput.png", "Figure 2. Achieved throughput plateaus around 49–51 requests/s even as virtual users increase five- to ten-fold.")
    add_metric_cards(doc, [("48.6", "req/s at 20 VUs", "pass"), ("48.8", "req/s at 100 VUs", "fail"), ("51.4", "req/s at 200 VUs", "fail"), ("3.48 s", "p95 at 200 VUs", "fail")])
    add_callout(doc, "Primary saturation signal", "Concurrency increased from 20 to 200 users, but throughput improved only 5.7% while p95 latency increased 14.8×. The dominant response is queue growth, consistent with the application’s 10-connection database pool and multi-query endpoints.", "fail")
    doc.add_heading("Working operating envelope", level=2)
    add_text(doc, "On this host and dataset, 20 concurrent read users is a defensible steady-state reference point: ~49 requests/s, 236 ms p95, and 0% failures. The baseline ramp is faster at lower concurrency but is not a sustained maximum. Capacity should be defined by the latency SLO, not by the absence of errors.")

    # Scale and DB.
    start_section(doc, "06", "Data scalability and database evidence", "Common search and deep offset pagination deteriorate as tenant data grows.")
    add_figure(doc, "dataset-scale.png", "Figure 3. Search/deep-page p95 increases from 18.8 ms to 1.23 s as the customer dataset grows from 1k to 50k rows.")
    add_table(
        doc,
        ["Dataset", "Customers", "Requests", "Req/s", "Average", "p95", "p99", "Threshold"],
        [
            ("Small", "1,000", "11,184", "62.1", "10.6 ms", "18.8 ms", "27.2 ms", "PASS"),
            ("Medium", "10,000", "7,831", "43.4", "78.8 ms", "210.2 ms", "291.0 ms", "PASS"),
            ("Large", "50,000", "2,789", "15.4", "495.6 ms", "1.23 s", "1.42 s", "FAIL"),
        ],
        widths=[2.0, 2.1, 2.1, 1.6, 2.0, 2.0, 2.0, 2.4],
        font_size=7.8,
    )
    doc.add_heading("EXPLAIN ANALYZE evidence", level=2)
    add_table(
        doc,
        ["Query", "Observed work at large scale", "Measured plan time"],
        [
            ("Vehicle common search", "Tenant index returns 25,000 vehicles; application expression and sort are evaluated before returning 20", "~61 ms"),
            ("Customer common search", "Tenant/phone index returns 50,000 customers; CONCAT_WS LIKE matches 5,000 and sorts to 20", "~145 ms"),
        ],
        widths=[3.4, 10.6, 2.8],
        font_size=8.0,
    )
    add_callout(doc, "Root cause", "The leading-wildcard CONCAT_WS search expression is not sargable, so existing tenant indexes cannot satisfy the search predicate. Concurrent data and COUNT queries, deep OFFSET pages, and dashboard fan-out amplify pressure on a 10-connection pool.", "fail")
    doc.add_heading("Recommended query changes", level=2)
    add_bullet(doc, "Replace common multi-column %term% scans with MySQL FULLTEXT or a purpose-built normalized search column and matching index.")
    add_bullet(doc, "Move list endpoints from deep OFFSET pagination to keyset/seek pagination using a stable (created_at, id) cursor.")
    add_bullet(doc, "Measure COUNT needs; return estimated totals or calculate counts separately where exact totals are not interaction-critical.")
    add_bullet(doc, "Consolidate dashboard aggregates or cache short-lived tenant summaries to reduce query fan-out.")

    # Race.
    start_section(doc, "07", "Transactional contention", "Uniqueness was preserved in both races; only the sale path translated losing attempts into controlled conflicts.")
    add_table(
        doc,
        ["Race", "Concurrency", "Successful transition", "Controlled conflict", "Unexpected error", "Result"],
        [
            ("Reserve one vehicle", "20", "1", "10 × HTTP 409", "9 × HTTP 500 deadlock", "FAIL"),
            ("Complete one sale", "20", "1", "19 × HTTP 409", "0", "PASS"),
        ],
        widths=[4.0, 2.0, 2.8, 3.3, 3.2, 1.7],
        font_size=8.0,
    )
    add_callout(doc, "Critical correctness finding", "The reservation race did not double-reserve the vehicle, but MySQL ER_LOCK_DEADLOCK escaped as HTTP 500 for 9 callers. This is a resilience defect: expected transactional contention must be retried or translated into a stable conflict/unavailable response.", "fail")
    doc.add_heading("Remediation design", level=2)
    add_number(doc, 1, "Acquire locks in a single, documented order across vehicle, reservation, offer, and related records.")
    add_number(doc, 2, "Retry MySQL deadlock/serialization errors (1213 / SQLSTATE 40001) with a small bounded exponential backoff and jitter.")
    add_number(doc, 3, "After the retry budget, return a controlled 409 conflict or 503 retryable response; never expose the database failure as a generic 500.")
    add_number(doc, 4, "Keep the existing uniqueness/state preconditions and add this 20-way test to the mandatory release gate.")

    # Resources and profiles.
    start_section(doc, "08", "Soak stability and runtime profiles", "Memory remained bounded during sustained average load; CPU samples and profiles point to query concurrency and password hashing as the main pressure sources.")
    add_figure(doc, "soak-memory.png", "Figure 4. RSS and heap remain bounded across the 30-minute soak. Short heap oscillations are consistent with garbage collection.")
    runtime = SUMMARY["runtime"]
    add_table(
        doc,
        ["Signal", "First 5 min", "Last 5 min", "Maximum / average", "Interpretation"],
        [
            ("RSS", f"{runtime['rssMb']['firstFiveMinuteAverage']:.2f} MB", f"{runtime['rssMb']['lastFiveMinuteAverage']:.2f} MB", f"{runtime['rssMb']['maximum']:.2f} MB max", "+0.44 MB; bounded"),
            ("Heap used", f"{runtime['heapUsedMb']['firstFiveMinuteAverage']:.2f} MB", f"{runtime['heapUsedMb']['lastFiveMinuteAverage']:.2f} MB", f"{runtime['heapUsedMb']['maximum']:.2f} MB max", "+0.01 MB; bounded"),
            ("Process CPU", "—", "—", f"{runtime['cpuPercent']['average']:.1f}% avg / {runtime['cpuPercent']['maximum']:.1f}% max", "Bursty saturation"),
            ("Event loop", "—", "—", f"p95 sample avg {runtime['eventLoopMs']['p95SampleAverage']:.1f} ms; max {runtime['eventLoopMs']['maximum']:.0f} ms", "Rare long stall observed"),
        ],
        widths=[2.4, 2.5, 2.5, 4.9, 4.5],
        font_size=7.8,
    )
    doc.add_heading("CPU and allocation profile", level=2)
    cpu_total = PROFILE["totalCpuProfileMilliseconds"]
    cpu_top = PROFILE["cpuTop"][:4]
    cpu_rows = []
    for entry in cpu_top:
        cpu_rows.append((entry["functionName"], Path(entry["url"]).name if entry["url"] else "runtime", f"{entry['milliseconds']:,.0f} ms", f"{entry['percent']:.2f}%"))
    add_table(doc, ["Hot function", "Module", "Sampled CPU", "Share"], cpu_rows, widths=[4.3, 6.2, 3.2, 2.3], font_size=8.0)
    add_text(doc, f"The CPU profile covers {cpu_total / 1000:.1f} seconds. Repeated bcryptjs _encipher frames dominate the hottest samples, aligning with the login-load collapse. Heap allocations include loader/LRU/MySQL buffers and a 0.50 MB loginTenantUser allocation; the soak data does not indicate a retained-memory leak.")
    add_callout(doc, "Authentication design implication", "Do not increase login concurrency blindly. First add rate limits and queue bounds, then evaluate native/worker-thread password verification or isolate authentication capacity. Any microservice result must use the same password cost and request schedule.", "warn")

    # Frontend.
    start_section(doc, "09", "Frontend quality audit", "Accessibility and best-practice scores are strong; performance and discoverability remain the main frontend opportunities.")
    add_figure(doc, "lighthouse-scores.png", "Figure 5. Lighthouse category scores for the public login page and authenticated dashboard.")
    add_table(
        doc,
        ["Page", "Performance", "Accessibility", "Best practices", "SEO"],
        [
            ("Login", str(LH["login"]["performance"]), str(LH["login"]["accessibility"]), str(LH["login"]["best-practices"]), str(LH["login"]["seo"])),
            ("Dashboard", str(LH["dashboard"]["performance"]), str(LH["dashboard"]["accessibility"]), str(LH["dashboard"]["best-practices"]), str(LH["dashboard"]["seo"])),
        ],
        widths=[4.8, 3.0, 3.0, 3.2, 2.8],
        font_size=8.5,
    )
    doc.add_heading("Priority frontend follow-up", level=2)
    add_bullet(doc, "Use the Lighthouse JSON/HTML artifacts to target the largest content and JavaScript opportunities on each page.")
    add_bullet(doc, "Compress and correctly size large visual assets; preload only truly critical resources.")
    add_bullet(doc, "Review route-level code splitting and defer non-critical dashboard bundles.")
    add_bullet(doc, "Add page-specific metadata where public discoverability matters; authenticated screens do not require aggressive SEO work.")
    add_callout(doc, "Interpretation", "Lighthouse ran locally against the production frontend. It is useful as a repeatable regression comparison, but field performance requires real-user monitoring or controlled remote infrastructure.", "info")

    # Priorities.
    start_section(doc, "10", "Prioritized engineering actions", "Fix correctness and backpressure first; then reduce database work and frontend transfer cost.")
    add_table(
        doc,
        ["Priority", "Action", "Why it matters", "Verification gate"],
        [
            ("P0", "Deadlock retry + deterministic lock order + stable response mapping", "Prevents expected contention from becoming HTTP 500", "20-way reservation: 1 success, 19 controlled conflicts"),
            ("P0", "Login rate limit, bounded queue, timeout, and password-hash capacity strategy", "Protects the event loop and database during auth bursts", "No dropped iterations at agreed login rate; p95 target met"),
            ("P1", "Sargable search (FULLTEXT/normalized index)", "Removes tenant-wide expression scans", "Large search p95 < 750 ms; plans avoid full tenant scan"),
            ("P1", "Keyset pagination and COUNT policy", "Prevents deep-page work from growing with offset", "Flat p95 from early to deep pages"),
            ("P1", "Pool/queue tuning with end-to-end timeouts and backpressure", "Makes overload explicit instead of silently queueing", "Throughput/latency knee documented; bounded queue"),
            ("P1", "Dashboard aggregate consolidation or short-lived cache", "Reduces per-request DB round trips", "Lower DB time and p95 without stale-data violation"),
            ("P2", "Asset optimization, route splitting, and metadata", "Raises Lighthouse performance/SEO", "Repeatable score and web-vitals improvement"),
        ],
        widths=[1.3, 5.2, 5.2, 5.1],
        font_size=7.5,
    )
    add_callout(doc, "Fair-comparison warning", "If these optimizations are applied only to the microservices version, the study will confound architecture with implementation changes. Either backport them to the monolith or report them explicitly as separate experimental factors.", "warn")

    # Microservices protocol.
    start_section(doc, "11", "Microservices comparison protocol", "Freeze the instrument now; change only the architecture treatment unless a change is explicitly documented.")
    doc.add_heading("Execution sequence", level=2)
    steps = [
        "Provision a clean comparison database and seed the same small, medium, and large datasets.",
        "Build production artifacts; record commit, runtime, database, service count, CPU/RAM limits, and network topology.",
        "Warm the system with the smoke profile, then run functional E2E and race tests.",
        "Run baseline, average, stress, spike, soak, breakpoint, auth, and search profiles in the same order with cooldowns.",
        "Capture per-service CPU, RSS, heap, event-loop, database, queue, and inter-service latency metrics.",
        "Repeat each profile 3–5 times; compare medians and dispersion against this monolith baseline.",
    ]
    for index, item in enumerate(steps, start=1):
        add_number(doc, index, item)
    doc.add_heading("Baseline comparison matrix", level=2)
    add_table(
        doc,
        ["Metric", "Monolith baseline", "Microservices result", "Preferred direction"],
        [
            ("E2E pass rate", "6/6 (100%)", "[enter]", "Equal"),
            ("Average-load throughput", "48.6 req/s", "[enter]", "Higher"),
            ("Average-load p95", "235.8 ms", "[enter]", "Lower"),
            ("Stress p95 at 100 VUs", "1.30 s", "[enter]", "Lower"),
            ("Breakpoint throughput at 200 VUs", "51.4 req/s", "[enter]", "Higher"),
            ("Spike p99", "22.39 s", "[enter]", "Lower"),
            ("Soak RSS drift", "+0.44 MB", "[enter total + per service]", "Bounded"),
            ("Large-data search p95", "1.23 s", "[enter]", "Lower"),
            ("Auth p95 / dropped", "40.43 s / 1,229", "[enter]", "Lower / zero"),
            ("Reservation race", "1 success / 10 conflicts / 9 errors", "[enter]", "1 / 19 / 0"),
            ("Sale race", "1 success / 19 conflicts / 0 errors", "[enter]", "Equal"),
        ],
        widths=[5.0, 4.0, 4.3, 3.5],
        font_size=7.7,
    )
    doc.add_heading("Additional microservices-only evidence", level=2)
    add_bullet(doc, "End-to-end latency broken down by gateway, service, database, message broker, and network hop.")
    add_bullet(doc, "Per-service scaling, queue depth, retry counts, circuit-breaker state, and distributed trace critical path.")
    add_bullet(doc, "Partial-failure tests: one dependency slow, unavailable, or returning errors; document recovery time and data consistency.")
    add_bullet(doc, "Operational cost proxy: total CPU-seconds, peak memory across all services, storage, and deployment complexity for the same workload.")

    # Appendix.
    start_section(doc, "A", "Appendix — endpoint and artifact index", "All results are retained as machine-readable evidence so tables and claims can be regenerated.")
    doc.add_heading("Performance endpoint matrix", level=2)
    add_table(
        doc,
        ["Workload", "Representative endpoints"],
        [
            ("Authentication", "POST /api/auth/login; GET /api/auth/me; refresh/logout through the browser session"),
            ("Mixed reads", "GET /api/tenant/dashboard, /vehicles, /customers, /leads, /test-drives, /offers, /reservations, /sales"),
            ("Search/scale", "GET /vehicles?search=…; /customers?search=…; deep pages for vehicles, customers, leads, and sales"),
            ("Reservation race", "POST /api/tenant/vehicles/:id/reservation"),
            ("Sale race", "PATCH /api/tenant/sales/:id/status { status: COMPLETED }"),
            ("Functional workflow", "UI and API calls spanning customer, vehicle, lead, activity, interest, test drive, offer, reservation, and sale"),
        ],
        widths=[4.0, 12.8],
        font_size=8.0,
    )
    doc.add_heading("Artifact map", level=2)
    add_table(
        doc,
        ["Evidence", "Path from repository root"],
        [
            ("Consolidated analysis", "testing/results/summary/monolith-summary.json"),
            ("System capture", "testing/results/summary/system.json"),
            ("CPU/heap analysis", "testing/results/summary/profile-summary.json and testing/results/profiles/"),
            ("k6 raw summaries", "testing/results/k6/*.json"),
            ("Playwright", "testing/results/e2e/results.json, html/, artifacts/"),
            ("Database plans", "testing/results/database/performance-audit.json"),
            ("Runtime telemetry", "testing/results/runtime/monolith-process.csv"),
            ("Lighthouse", "testing/results/lighthouse/*.json and *.html"),
            ("Charts", "testing/results/report-assets/*.png"),
            ("Reusable workloads", "testing/k6/*.js and frontend/tests/e2e/*.spec.ts"),
        ],
        widths=[5.0, 11.8],
        font_size=8.0,
    )
    doc.add_heading("Reproduction entry points", level=2)
    add_text(doc, "Backend: npm run build; npm run perf:seed -- medium|large; npm run perf:audit. Frontend: npm run build; npm run test:e2e; npm run lighthouse. Load: invoke the portable k6 executable with the selected script and environment variables documented in testing/k6/README.md.")
    add_callout(doc, "Traceability rule", "When this report is edited for the thesis, preserve the raw result filenames and execution metadata next to every derived table or chart. This separates measured evidence from interpretation and makes the microservices comparison auditable.", "info")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
